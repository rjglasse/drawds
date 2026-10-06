#!/usr/bin/env python3
"""Build a 16:9 slide deck (.pptx) from step images exported by the driver's `steps` command.

Each section is a folder the driver wrote (step PNGs plus steps.json). The deck has a title slide,
then per section a heading slide and one slide per step: the section and the step's header
("Step 2 of 3 · comparisons 1") at the top, the image fitted in the middle, the caption underneath
as editable text, and both in the speaker notes. Export the steps with captions=off (the caption is
on the slide instead) and background=off (transparent, on the slide's own background).

Google Slides: drop the .pptx into Google Drive (or File > Open > Upload in Slides); it converts.

    python3 deck.py out.pptx --title "Data structures, step by step" \\
        --section "Inserting 45 into a BST|Compare with each node on the way down|$TMPDIR/run-drawds/deck-bst"

Needs python-pptx and Pillow.
"""
import argparse
import json
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

WIDTH, HEIGHT = Inches(10), Inches(5.625)
MARGIN = Inches(0.45)
INK, GREY = RGBColor(0x1D, 0x1D, 0x1F), RGBColor(0x6E, 0x6E, 0x73)


def text(slide, x, y, w, h, value, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    frame = slide.shapes.add_textbox(x, y, w, h).text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color


def picture(slide, path, x, y, w, h):
    """The image as large as fits in the box, centred."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(path, Emu(x + (w - pw) // 2), Emu(y + (h - ph) // 2), Emu(pw), Emu(ph))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('out')
    parser.add_argument('--title', required=True)
    parser.add_argument('--subtitle', default='')
    parser.add_argument('--section', action='append', required=True, help='"title|description|folder"')
    args = parser.parse_args()

    deck = Presentation()
    deck.slide_width, deck.slide_height = WIDTH, HEIGHT
    blank = deck.slide_layouts[6]
    inner = WIDTH - 2 * MARGIN

    slide = deck.slides.add_slide(blank)
    text(slide, MARGIN, Inches(1.7), inner, Inches(1), args.title, 36, bold=True)
    text(slide, MARGIN, Inches(2.8), inner, Inches(1), args.subtitle, 18, GREY)

    for section in args.section:
        title, description, folder = section.split('|')
        steps = json.load(open(os.path.join(folder, 'steps.json')))
        slide = deck.slides.add_slide(blank)
        text(slide, MARGIN, Inches(1.9), inner, Inches(0.9), title, 30, bold=True)
        text(slide, MARGIN, Inches(2.8), inner, Inches(1), f'{description} ({len(steps)} steps)', 18, GREY)
        for step in steps:
            slide = deck.slides.add_slide(blank)
            text(slide, MARGIN, Inches(0.22), Inches(5.5), Inches(0.4), title, 14, GREY)
            text(slide, MARGIN + Inches(5.5), Inches(0.22), inner - Inches(5.5), Inches(0.4), step['header'], 14, GREY, align=PP_ALIGN.RIGHT)
            picture(slide, os.path.join(folder, step['file']), MARGIN, Inches(0.7), inner, Inches(3.85))
            text(slide, MARGIN, Inches(4.62), inner, Inches(0.85), step['caption'], 20, anchor=MSO_ANCHOR.TOP)
            slide.notes_slide.notes_text_frame.text = f"{step['header']}: {step['caption']}"

    deck.save(args.out)
    print(f'{args.out}: {len(deck.slides)} slides')


if __name__ == '__main__':
    main()
